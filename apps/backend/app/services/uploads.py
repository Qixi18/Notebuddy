"""Safe multi-format ingestion; one committed material and queued job per upload."""

from __future__ import annotations

import hashlib
import logging
import os
from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import BadZipFile, ZipFile

from sqlalchemy import func, select, update
from sqlalchemy.orm import Session

from app.core.config import settings
from app.db.models import (
    Asset,
    Course,
    Material,
    Notebook,
    NotebookAsset,
    ProcessingJob,
    new_id,
    utc_now,
)
from app.notes.assets import AssetUnavailable
from app.notes.notebook import bump_order_revision

logger = logging.getLogger(__name__)

DOCUMENT_MIME_TYPES = {
    ".pdf": "application/pdf",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class UploadRejected(ValueError):
    def __init__(self, message: str, status_code: int = 400) -> None:
        super().__init__(message)
        self.status_code = status_code


def _validate_container(path: Path, extension: str) -> None:
    if extension == ".pdf":
        with path.open("rb") as source:
            if source.read(5) != b"%PDF-":
                raise UploadRejected("文件扩展名是 PDF，但文件签名无效", 415)
        try:
            from pypdf import PdfReader
            page_count = len(PdfReader(str(path)).pages)
        except Exception as exc:
            raise UploadRejected("PDF 文件损坏或无法读取", 415) from exc
        if page_count > settings.max_pages:
            raise UploadRejected(f"PDF 页面数超过本地限制（最多 {settings.max_pages} 页）", 413)
        return
    try:
        with ZipFile(path) as archive:
            names = set(archive.namelist())
            required = {"ppt/presentation.xml"} if extension == ".pptx" else {"word/document.xml", "[Content_Types].xml"}
            if not required.issubset(names) or archive.testzip() is not None:
                raise UploadRejected(f"文件不是可读取的 {extension[1:].upper()}", 415)
        if extension == ".pptx":
            from pptx import Presentation

            page_count = len(Presentation(str(path)).slides)
            if page_count > settings.max_pages:
                raise UploadRejected(f"PPTX 页面数超过本地限制（最多 {settings.max_pages} 页）", 413)
        elif extension == ".docx":
            from docx import Document

            document = Document(str(path))
            position_count = len(document.paragraphs) + len(document.tables)
            if position_count > settings.max_pages:
                raise UploadRejected(f"DOCX 章节/位置数超过本地限制（最多 {settings.max_pages} 个）", 413)
    except UploadRejected:
        raise
    except BadZipFile as exc:
        raise UploadRejected(f"文件不是可读取的 {extension[1:].upper()}", 415) from exc
    except Exception as exc:
        raise UploadRejected(f"文件不是可读取的 {extension[1:].upper()}", 415) from exc


def validate_upload_container(path: Path, extension: str) -> None:
    """Public validation entry for non-material uploads."""
    _validate_container(path, extension)


def _safe_data_path(key: str) -> Path:
    root = settings.data_dir.resolve()
    candidate = root / key
    try:
        path = candidate.resolve()
    except (OSError, RuntimeError) as exc:
        raise AssetUnavailable("资源存储路径不可访问") from exc
    if not path.is_relative_to(root) or path == root:
        raise AssetUnavailable("资源存储路径无效")
    current = candidate
    while current != root:
        if current.is_symlink():
            raise AssetUnavailable("资源存储路径包含符号链接")
        current = current.parent
    return candidate


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while chunk := stream.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def _fsync_directory(path: Path) -> None:
    try:
        descriptor = os.open(path, os.O_RDONLY)
    except OSError:
        return
    try:
        os.fsync(descriptor)
    except OSError:
        pass
    finally:
        os.close(descriptor)


def reconcile_staged_document_assets(db: Session, *, limit: int = 50) -> int:
    """Finish or quarantine document uploads left between SQLite and rename."""
    assets = db.scalars(select(Asset).where(
        Asset.state == "staging", Asset.media_type == "document",
    ).order_by(Asset.created_at).limit(limit)).all()
    reconciled = 0
    for asset in assets:
        try:
            final = _safe_data_path(asset.storage_key)
            stage = _safe_data_path(f"staging/{asset.id}.upload")
            final.parent.mkdir(parents=True, exist_ok=True)
            final_digest = _sha256_file(final) if final.is_file() else None
            stage_digest = _sha256_file(stage) if stage.is_file() else None
            if final_digest == asset.sha256:
                digest = final_digest
                stage.unlink(missing_ok=True)
            elif stage_digest == asset.sha256:
                # A previous promotion may have left a corrupt destination;
                # preserve the verified staged copy until it replaces it.
                os.replace(stage, final)
                _fsync_directory(final.parent)
                digest = stage_digest
            elif final_digest is not None:
                digest = final_digest
            elif stage_digest is not None:
                digest = stage_digest
                os.replace(stage, final)
                _fsync_directory(final.parent)
            else:
                digest = ""
        except (OSError, AssetUnavailable):
            logger.exception("Could not reconcile staged document Asset %s", asset.id)
            continue

        material = db.scalar(select(Material).where(Material.asset_id == asset.id))
        if not digest:
            asset.state = "missing" if not final.exists() and not stage.exists() else "quarantined"
            if material is not None:
                material.status = "failed"
                material.document_warning = "上传文件缺失或哈希校验失败；原始资料未对外开放"
                job = db.scalar(select(ProcessingJob).where(ProcessingJob.material_id == material.id))
                if job is not None and job.status in {"pending", "processing"}:
                    job.status = "failed"
                    job.phase = "failed"
                    job.error_code = "asset_integrity_failed"
                    job.error_message = material.document_warning
                    job.finished_at = utc_now()
                    job.updated_at = job.finished_at
            db.commit()
            reconciled += 1
            continue
        if digest != asset.sha256:
            asset.state = "quarantined"
            if material is not None:
                material.status = "failed"
                material.document_warning = "上传文件 SHA-256 校验失败，已隔离"
            job = db.scalar(select(ProcessingJob).where(ProcessingJob.material_id == material.id)) if material else None
            if job is not None and job.status in {"pending", "processing"}:
                job.status = "failed"
                job.phase = "failed"
                job.error_code = "asset_hash_mismatch"
                job.error_message = material.document_warning
                job.finished_at = utc_now()
                job.updated_at = job.finished_at
            db.commit()
            reconciled += 1
            continue
        asset.state = "ready"
        asset.verified_at = utc_now()
        db.commit()
        reconciled += 1
    return reconciled


def ingest_document(
    db: Session, *, course_id: str | None, source: object, filename: str,
    lecture_title: str, topic_title: str, media_type: str | None,
    allow_duplicate: bool = False, idempotency_key: str | None = None,
    notebook_id: str | None = None,
) -> tuple[Material, ProcessingJob]:
    extension = Path(filename).suffix.lower()
    if extension not in settings.allowed_extensions:
        raise UploadRejected("当前支持 PPTX、PDF 和 DOCX 文件", 415)
    if notebook_id is None:
        raise UploadRejected("请明确选择要归属的笔记本；不会自动选择课程默认笔记", 422)
    book = db.get(Notebook, notebook_id)
    if book is None or book.deleted_at is not None:
        raise UploadRejected("笔记本不存在", 404)
    if course_id is not None:
        course = db.get(Course, course_id)
        if course is None or course.deleted_at is not None:
            raise UploadRejected("课程不存在", 404)
        if book.course_id != course_id:
            raise UploadRejected("所选笔记不属于该课程", 409)
    notebook_id = book.id
    key = hashlib.sha256(f"{notebook_id}:{idempotency_key}".encode()).hexdigest() if idempotency_key else None
    if key:
        prior = db.scalar(select(ProcessingJob).where(ProcessingJob.idempotency_key == key))
        if prior is not None:
            if prior.material.notebook_id != notebook_id:
                raise UploadRejected("重复请求不属于当前笔记", 409)
            return prior.material, prior

    settings.data_dir.joinpath("staging").mkdir(parents=True, exist_ok=True)
    settings.originals_dir.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256()
    size = 0
    temporary = None
    destination = None
    staging_path = None
    committed = False
    try:
        with NamedTemporaryFile(dir=settings.data_dir / "staging", prefix="upload-", suffix=".tmp", delete=False) as output:
            temporary = Path(output.name)
            while chunk := source.read(1024 * 1024):
                size += len(chunk)
                if size > settings.max_upload_bytes:
                    raise UploadRejected("文件超过本地配置的大小限制", 413)
                output.write(chunk)
                digest.update(chunk)
            output.flush()
            os.fsync(output.fileno())
        if size == 0:
            raise UploadRejected("文件为空或无法读取", 415)
        _validate_container(temporary, extension)
        content_hash = digest.hexdigest()
        # Recheck after taking the SQLite write lock: concurrent retries may
        # both have passed the early lookup while streaming their files.
        db.execute(update(Notebook).where(Notebook.id == notebook_id).values(
            order_revision=Notebook.order_revision))
        if key:
            prior = db.scalar(select(ProcessingJob).where(ProcessingJob.idempotency_key == key))
            if prior is not None:
                material = prior.material
                db.commit()
                return material, prior
        duplicate = db.scalar(select(Material).where(
            Material.notebook_id == notebook_id, Material.content_hash == content_hash,
            Material.deleted_at.is_(None),
        ))
        if duplicate is not None and not allow_duplicate:
            raise UploadRejected(f"该笔记已上传相同文件：{duplicate.lecture_title}。如需作为新讲次导入，请确认。", 409)
        material_id = new_id()
        # Serialize chapter allocation with reorders and other uploads.
        bump_order_revision(db, book.course_id, notebook_id)
        last_order = db.scalar(select(func.max(Material.chapter_order)).where(Material.notebook_id == notebook_id))
        destination = settings.originals_dir / f"{material_id}{extension}"
        staging_path = settings.data_dir / "staging" / f"{material_id}.upload"
        os.replace(temporary, staging_path)
        temporary = None
        _fsync_directory(staging_path.parent)
        created_at = utc_now()
        storage_key = f"originals/{destination.name}"
        asset = Asset(
            id=material_id, source_type="user_upload", media_type="document",
            mime_type=DOCUMENT_MIME_TYPES[extension], extension=extension,
            original_name=Path(filename).name[:255], storage_key=storage_key,
            byte_size=size, sha256=content_hash, state="staging",
            created_at=created_at, verified_at=None,
        )
        material = Material(
            id=material_id, asset_id=material_id, course_id=None, notebook_id=notebook_id,
            lecture_title=lecture_title.strip()[:200] or filename[:200],
            topic_title=topic_title.strip()[:200] or None,
            original_filename=Path(filename).name[:255], stored_filename=destination.name,
            media_type=media_type, size_bytes=size, content_hash=content_hash,
            status="pending",
            chapter_order=0 if last_order is None else last_order + 1,
        )
        job = ProcessingJob(material_id=material_id, kind="parse_material", idempotency_key=key)
        db.add_all([
            asset,
            NotebookAsset(notebook_id=notebook_id, asset_id=material_id,
                          role="source_material", is_owner=True, created_at=created_at),
            material, job,
        ])
        db.commit()
        committed = True
        try:
            os.replace(staging_path, destination)
            staging_path = None
            asset.state = "ready"
            asset.verified_at = utc_now()
            db.commit()
        except OSError:
            # Keep the durable staging row and job. Worker reconciliation can
            # finish the rename after the filesystem recovers.
            db.rollback()
            logger.exception("Could not finalize staged document Asset %s", material_id)
        except Exception:
            # The final rename may have succeeded while SQLite commit failed;
            # the reconciler checks both paths and the stored SHA-256.
            db.rollback()
            logger.exception("Could not mark document Asset %s ready", material_id)
        db.refresh(material)
        db.refresh(job)
        return material, job
    except Exception:
        if not committed:
            db.rollback()
            if staging_path is not None:
                staging_path.unlink(missing_ok=True)
            if destination is not None:
                destination.unlink(missing_ok=True)
        raise
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def ingest_pptx(db: Session, **kwargs):
    """Backward-compatible name retained for integrations from phase one."""
    return ingest_document(db, **kwargs)
