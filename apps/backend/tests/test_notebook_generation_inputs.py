from __future__ import annotations

import json
from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, datetime
from io import BytesIO
from threading import Barrier
from types import SimpleNamespace

import pytest
from docx import Document
from sqlalchemy import create_engine, event, select
from sqlalchemy.orm import Session

from app.db.database import Base
from app.db.model_utils import new_id
from app.db.models import (
    Asset,
    Material,
    MaterialPage,
    Notebook,
    NotebookAsset,
    NotebookGenerationInput,
    NotebookGenerationJob,
    NotebookUploadBatch,
    ProcessingJob,
)
from app.services.notebook_generation_inputs import (
    close_upload_batch,
    coordinate_generation_inputs,
    register_generation_input,
)

pytest_plugins = ["tests.test_notebook_api"]


@pytest.fixture
def db():
    engine = create_engine("sqlite:///:memory:")
    with engine.begin() as connection:
        connection.exec_driver_sql("PRAGMA foreign_keys=ON")
    Base.metadata.create_all(engine)
    with Session(engine, autoflush=False) as session:
        yield session
    engine.dispose()


def _parsed_material(db: Session):
    now = datetime.now(UTC)
    notebook = Notebook(title="自动生成", tree_slot=0, document_json='{"type":"doc","content":[]}')
    batch = NotebookUploadBatch(id=new_id(), notebook_id="pending", created_at=now, updated_at=now)
    asset = Asset(
        source_type="user_upload", media_type="document", mime_type="application/pdf",
        extension=".pdf", original_name="lesson.pdf", storage_key=f"assets/{new_id()}",
        byte_size=10, sha256="a" * 64, state="ready", created_at=now,
    )
    db.add(notebook)
    db.flush()
    batch.notebook_id = notebook.id
    db.add_all([batch, asset])
    db.flush()
    asset.upload_batch_id = batch.id
    db.add(NotebookAsset(
        notebook_id=notebook.id, asset_id=asset.id, role="source_material", is_owner=True,
    ))
    material = Material(
        notebook_id=notebook.id, asset_id=asset.id, lecture_title="lesson",
        original_filename="lesson.pdf", stored_filename="lesson.pdf", size_bytes=10,
        status="completed", content_hash="a" * 64, parser_version="test-v1",
    )
    db.add(material)
    db.flush()
    parse_job = ProcessingJob(
        material_id=material.id, kind="parse_material", status="completed",
        finished_at=now,
    )
    db.add(parse_job)
    db.add(MaterialPage(
        material_id=material.id, page_number=1, raw_text="新增知识点",
        is_active=True,
    ))
    db.flush()
    register_generation_input(
        db, notebook_id=notebook.id, asset_id=asset.id, batch_id=batch.id,
        parse_kind="material", parse_job_id=parse_job.id,
    )
    batch.closed_at = now
    db.commit()
    return notebook, batch, asset


def test_coordinator_enqueues_closed_batch_once_and_keeps_asset_snapshot(db, monkeypatch):
    notebook, _batch, asset = _parsed_material(db)
    calls: list[tuple[str, list[str], str]] = []

    def create_job(notebook_id, payload, key, session):
        calls.append((notebook_id, payload.asset_ids, key))
        job = NotebookGenerationJob(
            id=new_id(), notebook_id=notebook_id, idempotency_key=key,
            request_fingerprint="b" * 64, status="pending", base_document_revision=0,
            input_snapshot_json=json.dumps({"sources": [{"asset_id": asset.id}]}),
            model="test", prompt_version="test", schema_version=2, attempt_count=0,
        )
        session.add(job)
        session.flush()
        return {"id": job.id, "status": job.status}

    monkeypatch.setattr("app.api.v1.generation.create_generation_job", create_job)
    assert coordinate_generation_inputs(db) == 1
    row = db.scalar(select(NotebookGenerationInput).where(
        NotebookGenerationInput.asset_id == asset.id,
    ))
    assert row is not None and row.status == "generating"
    assert row.generation_job_id is not None
    assert calls == [(notebook.id, [asset.id], calls[0][2])]
    assert coordinate_generation_inputs(db) == 0
    assert len(calls) == 1


def test_completed_asset_cannot_be_reopened_by_reparse_or_overlapping_batch(db):
    notebook, _batch, asset = _parsed_material(db)
    row = db.scalar(select(NotebookGenerationInput).where(
        NotebookGenerationInput.asset_id == asset.id,
    ))
    assert row is not None
    row.status = "completed"
    old_batch_id = row.batch_id
    next_batch = NotebookUploadBatch(
        id=new_id(), notebook_id=notebook.id,
        created_at=datetime.now(UTC), updated_at=datetime.now(UTC),
    )
    db.add(next_batch)
    db.flush()
    register_generation_input(
        db, notebook_id=notebook.id, asset_id=asset.id, batch_id=next_batch.id,
        parse_kind="material", parse_job_id=row.parse_job_id,
    )
    assert row.status == "completed"
    assert row.batch_id == old_batch_id


def test_batch_status_can_be_closed_without_dependent_browser_callback(db):
    notebook, batch, _ = _parsed_material(db)
    assert batch.closed_at is not None
    assert close_upload_batch(db, notebook.id, batch.id).closed_at is not None


def test_parse_failure_is_visible_and_recovered_parse_is_enqueued(db, monkeypatch):
    _notebook, _batch, asset = _parsed_material(db)
    row = db.scalar(select(NotebookGenerationInput).where(
        NotebookGenerationInput.asset_id == asset.id,
    ))
    material = db.scalar(select(Material).where(Material.asset_id == asset.id))
    parse_job = db.scalar(select(ProcessingJob).where(ProcessingJob.material_id == material.id))
    assert row is not None and material is not None and parse_job is not None
    parse_job.status = "failed"
    material.status = "failed"
    db.commit()
    assert coordinate_generation_inputs(db) == 0
    assert row.status == "parse_failed"
    assert row.error_code == "parse_failed"

    parse_job.status = "completed"
    material.status = "completed"
    row.status = "waiting_parse"  # The retry endpoint persists this transition.
    db.commit()

    def create_job(notebook_id, payload, key, session):
        job = NotebookGenerationJob(
            id=new_id(), notebook_id=notebook_id, idempotency_key=key,
            request_fingerprint="d" * 64, status="pending", base_document_revision=0,
            input_snapshot_json=json.dumps({"sources": [{"asset_id": asset.id}]}),
            model="test", prompt_version="test", schema_version=2, attempt_count=0,
        )
        session.add(job)
        session.flush()
        return {"id": job.id, "status": job.status}

    monkeypatch.setattr("app.api.v1.generation.create_generation_job", create_job)
    assert coordinate_generation_inputs(db) == 1
    assert row.status == "generating"


def test_notebook_upload_persists_batch_asset_parse_and_generation_input(
    setup, monkeypatch, tmp_path,
):
    client, db, course, *_ = setup
    book_response = client.post(
        f"/api/v1/courses/{course.id}/notebooks", json={"title": "上传自动编排"},
    )
    assert book_response.status_code == 201, book_response.text
    notebook_id = book_response.json()["id"]
    monkeypatch.setattr("app.services.uploads.settings", SimpleNamespace(
        data_dir=tmp_path,
        originals_dir=tmp_path / "originals",
        max_upload_bytes=10 * 1024 * 1024,
            max_pages=100,
        allowed_extensions=(".docx",),
    ))
    source = BytesIO()
    document = Document()
    document.add_paragraph("新上传的学习资料")
    document.save(source)
    batch_id = new_id()
    response = client.post(
        f"/api/v1/notebooks/{notebook_id}/files",
        headers={"X-Upload-Batch-ID": batch_id},
        files={"file": ("lesson.docx", source.getvalue(),
                         "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 201, response.text
    payload = response.json()
    assert payload["upload_batch_id"] == batch_id
    material = db.get(Material, payload["material"]["id"])
    assert material is not None and material.asset_id is not None
    asset = db.get(Asset, material.asset_id)
    row = db.scalar(select(NotebookGenerationInput).where(
        NotebookGenerationInput.asset_id == material.asset_id,
    ))
    batch = db.get(NotebookUploadBatch, batch_id)
    assert asset is not None and asset.upload_batch_id == batch_id and asset.sha256
    assert row is not None and row.parse_job_id == payload["job"]["id"]
    assert row.status == "waiting_parse" and batch is not None and batch.closed_at is None


def test_upload_parse_runner_and_generation_worker_append_official_document(
    setup, monkeypatch, tmp_path,
):
    """Exercise the durable path without a browser callback or generation API call."""
    from sqlalchemy.orm import sessionmaker

    from app.core.config import settings
    from app.db.models import NotebookDocumentRevision
    from app.workers import material_worker, notebook_generation_worker, runner

    client, db, course, *_ = setup
    notebook_response = client.post(
        f"/api/v1/courses/{course.id}/notebooks", json={"title": "端到端自动生成"},
    )
    assert notebook_response.status_code == 201, notebook_response.text
    notebook_id = notebook_response.json()["id"]
    document_url = f"/api/v1/courses/{course.id}/notebooks/{notebook_id}/document"
    original_node = {"type": "heading", "attrs": {"level": 1, "blockId": "user-original"},
                     "content": [{"type": "text", "text": "用户原有内容"}]}
    initial_save = client.put(document_url, json={
        "schema_version": 2, "content": {"type": "doc", "content": [original_node]},
        "expected_revision": 0,
    })
    assert initial_save.status_code == 200, initial_save.text
    monkeypatch.setattr(settings, "data_dir", tmp_path)
    monkeypatch.setattr(settings, "max_upload_bytes", 10 * 1024 * 1024)
    monkeypatch.setattr(settings, "max_pages", 100)
    monkeypatch.setattr(settings, "allowed_extensions", (".docx",))

    factory = sessionmaker(bind=db.get_bind(), autoflush=False, expire_on_commit=False)
    monkeypatch.setattr(material_worker, "SessionLocal", factory)
    monkeypatch.setattr(runner, "SessionLocal", factory)
    monkeypatch.setattr(notebook_generation_worker, "SessionLocal", factory)
    monkeypatch.setattr(runner, "claim_next_asset_job", lambda: None)
    monkeypatch.setattr(runner, "claim_next_forest_job", lambda: None)
    monkeypatch.setattr(material_worker, "extract_material_knowledge", lambda *_args: None)
    monkeypatch.setattr(material_worker, "index_material", lambda *_args: None)
    monkeypatch.setattr(material_worker, "enqueue_relations", lambda *_args, **_kwargs: None)

    batch_id = new_id()
    uploaded = []
    for title, text in (("第一份资料", "矩阵的秩表示线性独立信息的数量。"),
                        ("第二份资料", "秩与线性方程组解的结构密切相关。")):
        source = BytesIO()
        document = Document()
        document.add_paragraph(text)
        document.save(source)
        response = client.post(
            f"/api/v1/notebooks/{notebook_id}/files",
            headers={"X-Upload-Batch-ID": batch_id},
            files={"file": (f"{title}.docx", source.getvalue(),
                             "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            data={"lecture_title": title},
        )
        assert response.status_code == 201, response.text
        uploaded.append(response.json())
    assert {item["upload_batch_id"] for item in uploaded} == {batch_id}

    # Execute the real parser queue claims and parser Worker for both uploaded files.
    for item in uploaded:
        job_id = runner.claim_next_job()
        assert job_id == item["job"]["id"]
        material_worker.process_material(job_id)
    assert all(db.get(ProcessingJob, item["job"]["id"]).status == "completed" for item in uploaded)

    # Simulate the page closing: no close-batch API or generation endpoint call.
    batch = db.get(NotebookUploadBatch, batch_id)
    batch.updated_at = datetime.now(UTC).replace(year=datetime.now(UTC).year - 1)
    db.commit()
    output = {"type": "doc", "content": [
        {"type": "heading", "attrs": {"level": 2},
         "content": [{"type": "text", "text": "矩阵的秩"}]},
        {"type": "paragraph", "content": [
            {"type": "text", "text": "秩刻画了矩阵中线性独立信息的数量，并影响方程组的解。"},
        ]},
    ]}
    def save_during_model_call(*_args, **_kwargs):
        # A user save lands after generation starts and before the Worker reads
        # the latest document for its append transaction.
        newer_content = {"type": "doc", "content": [original_node, {
            "type": "paragraph", "attrs": {"blockId": "user-during-generation"},
            "content": [{"type": "text", "text": "生成期间用户新增的段落"}],
        }]}
        response = client.put(document_url, json={
            "schema_version": 2, "content": newer_content, "expected_revision": 1,
        })
        assert response.status_code == 200, response.text
        return output

    monkeypatch.setattr(
        "app.workers.notebook_generation_worker.DeepSeekClient.complete_json",
        save_during_model_call,
    )
    monkeypatch.setattr(runner.time, "monotonic", lambda: 10_000)
    monkeypatch.setattr(runner, "_last_generation_coordination", 0)
    monkeypatch.setattr(runner, "_jobs_since_generation", 0)

    # Crash window 1: all parse jobs committed, but the process dies before the
    # durable scan can create the generation job. A fresh scan must recover it.
    coordinate = runner.coordinate_generation_inputs
    monkeypatch.setattr(runner, "coordinate_generation_inputs", lambda _db: (_ for _ in ()).throw(
        RuntimeError("simulated runner interruption before enqueue"),
    ))
    assert runner.run_once() is False
    assert db.scalar(select(NotebookGenerationJob.id).where(
        NotebookGenerationJob.notebook_id == notebook_id,
    )) is None

    # Crash window 2: coordinator committed the generation job, then the process
    # stopped before claiming it. The next Runner instance consumes that row.
    monkeypatch.setattr(runner, "coordinate_generation_inputs", coordinate)
    claim_generation = runner.claim_next_generation_job
    monkeypatch.setattr(runner, "claim_next_generation_job", lambda: None)
    monkeypatch.setattr(runner, "_last_generation_coordination", 0)
    assert runner.run_once() is False
    pending = db.scalar(select(NotebookGenerationJob).where(
        NotebookGenerationJob.notebook_id == notebook_id,
    ))
    assert pending is not None and pending.status == "pending"
    monkeypatch.setattr(runner, "claim_next_generation_job", claim_generation)

    # The restarted Runner claims the durable job and dispatches the real Worker.
    assert runner.run_once() is True
    db.expire_all()
    persisted = db.get(Notebook, notebook_id)
    document_json = json.loads(persisted.document_json)["doc"]
    assert [node["type"] for node in document_json["content"]] == [
        "heading", "paragraph", "heading", "paragraph",
    ]
    assert document_json["content"][0] == original_node
    assert "生成期间用户新增的段落" in json.dumps(document_json, ensure_ascii=False)
    assert persisted.document_revision == 3
    rows = db.scalars(select(NotebookGenerationInput).where(
        NotebookGenerationInput.batch_id == batch_id,
    )).all()
    assert len(rows) == 2 and all(row.status == "completed" for row in rows)
    jobs = db.scalars(select(NotebookGenerationJob).where(
        NotebookGenerationJob.notebook_id == notebook_id,
    )).all()
    assert len(jobs) == 1 and jobs[0].status == "completed"
    assert all(row.generation_job_id == jobs[0].id for row in rows)
    snapshots = db.scalars(select(NotebookDocumentRevision).where(
        NotebookDocumentRevision.notebook_id == notebook_id,
    ).order_by(NotebookDocumentRevision.revision)).all()
    assert [snapshot.revision for snapshot in snapshots] == [1, 2, 3]
    assert snapshots[-1].origin == "ai_append"

    # At-least-once Runner/Worker delivery after the committed result is harmless.
    notebook_generation_worker.process_generation_job(jobs[0].id)
    assert runner.run_once() is False
    db.expire_all()
    assert db.get(Notebook, notebook_id).document_revision == 3


def test_partial_parse_failure_then_retry_generates_only_recovered_input(
    setup, monkeypatch, tmp_path,
):
    from sqlalchemy.orm import sessionmaker

    from app.core.config import settings
    from app.workers import material_worker, notebook_generation_worker, runner

    client, db, course, *_ = setup
    notebook_id = client.post(
        f"/api/v1/courses/{course.id}/notebooks", json={"title": "部分解析恢复"},
    ).json()["id"]
    monkeypatch.setattr(settings, "data_dir", tmp_path)
    monkeypatch.setattr(settings, "max_upload_bytes", 10 * 1024 * 1024)
    monkeypatch.setattr(settings, "max_pages", 100)
    monkeypatch.setattr(settings, "allowed_extensions", (".docx",))
    factory = sessionmaker(bind=db.get_bind(), autoflush=False, expire_on_commit=False)
    monkeypatch.setattr(material_worker, "SessionLocal", factory)
    monkeypatch.setattr(runner, "SessionLocal", factory)
    monkeypatch.setattr(notebook_generation_worker, "SessionLocal", factory)
    monkeypatch.setattr(runner, "claim_next_asset_job", lambda: None)
    monkeypatch.setattr(runner, "claim_next_forest_job", lambda: None)
    monkeypatch.setattr(material_worker, "extract_material_knowledge", lambda *_args: None)
    monkeypatch.setattr(material_worker, "index_material", lambda *_args: None)
    monkeypatch.setattr(material_worker, "enqueue_relations", lambda *_args, **_kwargs: None)

    batch_id = new_id()
    uploaded = []
    for title, text in (("待恢复", "这份资料首次解析失败后可以恢复。"),
                        ("成功资料", "这份资料应当先独立生成。")):
        source = BytesIO()
        document = Document()
        document.add_paragraph(text)
        document.save(source)
        response = client.post(
            f"/api/v1/notebooks/{notebook_id}/files",
            headers={"X-Upload-Batch-ID": batch_id},
            files={"file": (f"{title}.docx", source.getvalue(),
                             "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        assert response.status_code == 201, response.text
        uploaded.append(response.json())

    failed_material_id = uploaded[0]["material"]["id"]
    parse_document = material_worker.parse_document

    def fail_one_material_once(path):
        if path.name == db.get(Material, failed_material_id).stored_filename:
            raise ValueError("simulated parse failure")
        return parse_document(path)

    monkeypatch.setattr(material_worker, "parse_document", fail_one_material_once)
    for item in uploaded:
        job_id = runner.claim_next_job()
        assert job_id == item["job"]["id"]
        material_worker.process_material(job_id)
    failed_row = db.scalar(select(NotebookGenerationInput).where(
        NotebookGenerationInput.asset_id == failed_material_id,
    ))
    assert failed_row.status == "waiting_parse"  # Coordinator records terminal status.
    monkeypatch.setattr(material_worker, "parse_document", parse_document)

    batch = db.get(NotebookUploadBatch, batch_id)
    batch.updated_at = datetime.now(UTC).replace(year=datetime.now(UTC).year - 1)
    db.commit()
    generated_texts = iter(("先生成成功资料", "再补充恢复资料"))
    monkeypatch.setattr(
        "app.workers.notebook_generation_worker.DeepSeekClient.complete_json",
        lambda *_args, **_kwargs: {"type": "doc", "content": [{
            "type": "paragraph", "content": [{"type": "text", "text": next(generated_texts)}],
        }]},
    )
    monkeypatch.setattr(runner.time, "monotonic", lambda: 20_000)
    monkeypatch.setattr(runner, "_last_generation_coordination", 0)
    monkeypatch.setattr(runner, "_jobs_since_generation", 0)

    # The successful member proceeds while the failed file remains recoverable.
    assert runner.run_once() is True
    db.expire_all()
    failed_row = db.scalar(select(NotebookGenerationInput).where(
        NotebookGenerationInput.asset_id == failed_material_id,
    ))
    assert failed_row.status == "parse_failed"
    status = client.get(
        f"/api/v1/notebooks/{notebook_id}/upload-batches/{batch_id}",
    ).json()
    assert next(item for item in status["inputs"] if item["asset_id"] == failed_material_id)["status"] == "parse_failed"
    successful_asset_id = uploaded[1]["material"]["asset_id"]
    successful_row = db.scalar(select(NotebookGenerationInput).where(
        NotebookGenerationInput.asset_id == successful_asset_id,
    ))
    assert successful_row.status == "completed"

    retry_response = client.post(f"/api/v1/materials/{failed_material_id}/retry")
    assert retry_response.status_code == 202, retry_response.text
    retry_job_id = runner.claim_next_job()
    assert retry_job_id == retry_response.json()["id"]
    material_worker.process_material(retry_job_id)
    monkeypatch.setattr(runner, "_last_generation_coordination", 0)
    assert runner.run_once() is True
    db.expire_all()
    assert db.get(Notebook, notebook_id).document_revision == 2
    assert successful_row.status == "completed"
    assert failed_row.status == "completed"
    jobs = db.scalars(select(NotebookGenerationJob).where(
        NotebookGenerationJob.notebook_id == notebook_id,
    ).order_by(NotebookGenerationJob.created_at)).all()
    assert len(jobs) == 2
    recovered_snapshot = json.loads(jobs[-1].input_snapshot_json)
    assert [source["asset_id"] for source in recovered_snapshot["sources"]] == [failed_material_id]


def test_two_workers_racing_for_one_generation_job_append_exactly_once(tmp_path, monkeypatch):
    """Two independent Worker sessions race at the real durable claim boundary."""
    from sqlalchemy.orm import sessionmaker

    from app.db.models import NotebookDocumentRevision
    from app.notes.document import load_document
    from app.workers import notebook_generation_worker

    engine = create_engine(
        f"sqlite:///{(tmp_path / 'worker-race.sqlite3').as_posix()}",
        connect_args={"timeout": 10},
    )
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)
    with factory() as db:
        notebook, _batch, asset = _parsed_material(db)
        job = NotebookGenerationJob(
            id=new_id(), notebook_id=notebook.id, idempotency_key="race-once",
            request_fingerprint="e" * 64, status="pending", base_document_revision=0,
            input_snapshot_json=json.dumps({
                "sources": [{"asset_id": asset.id, "text": "竞态测试资料"}],
                "current_document": {"type": "doc", "content": []},
            }),
            model="test", prompt_version="test", schema_version=2, attempt_count=0,
        )
        db.add(job)
        db.flush()
        row = db.get(NotebookGenerationInput, db.scalar(select(NotebookGenerationInput.id).where(
            NotebookGenerationInput.asset_id == asset.id,
        )))
        row.status = "generating"
        row.generation_job_id = job.id
        notebook_id, asset_id, job_id = notebook.id, asset.id, job.id
        db.commit()

    barrier = Barrier(2)

    @event.listens_for(engine, "before_cursor_execute")
    def synchronize_claims(_conn, _cursor, statement, _parameters, _context, _many):
        if ("FROM notebook_generation_jobs" in statement
                and "notebook_generation_jobs.status = ?" in statement
                and "LIMIT" in statement):
            barrier.wait(timeout=10)

    monkeypatch.setattr(notebook_generation_worker, "SessionLocal", factory)
    with ThreadPoolExecutor(max_workers=2) as pool:
        claims = list(pool.map(lambda _worker: notebook_generation_worker.claim_next_generation_job(), range(2)))

    assert sorted(claims, key=lambda value: value is None) == [job_id, None]

    monkeypatch.setattr(
        "app.workers.notebook_generation_worker.DeepSeekClient.complete_json",
        lambda *_args, **_kwargs: {"type": "doc", "content": [{
            "type": "paragraph", "content":[{"type": "text", "text": "只追加一次"}],
        }]},
    )
    winner = next(claim for claim in claims if claim is not None)
    notebook_generation_worker.process_generation_job(winner)

    with factory() as db:
        notebook = db.get(Notebook, notebook_id)
        document = load_document(notebook.document_json).content
        node_text = json.dumps(document, ensure_ascii=False)
        generation = db.get(NotebookGenerationJob, job_id)
        input_row = db.scalar(select(NotebookGenerationInput).where(
            NotebookGenerationInput.asset_id == asset_id,
        ))
        revisions = db.scalars(select(NotebookDocumentRevision).where(
            NotebookDocumentRevision.notebook_id == notebook_id,
        )).all()
        assert node_text.count("只追加一次") == 1
        assert notebook.document_revision == 1
        assert len(revisions) == 1 and revisions[0].origin == "ai_append"
        assert generation.status == "completed" and generation.attempt_count == 1
        assert input_row.status == "completed" and input_row.generation_job_id == job_id
        assert db.query(NotebookGenerationJob).filter_by(notebook_id=notebook_id).count() == 1
    engine.dispose()


def test_restore_history_then_upload_new_asset_appends_without_reprocessing_old_inputs(
    setup, monkeypatch, tmp_path,
):
    from sqlalchemy.orm import sessionmaker

    from app.core.config import settings
    from app.db.models import NotebookDocumentRevision
    from app.workers import material_worker, notebook_generation_worker, runner

    client, db, course, *_ = setup
    notebook_id = client.post(
        f"/api/v1/courses/{course.id}/notebooks", json={"title": "恢复后继续追加"},
    ).json()["id"]
    document_url = f"/api/v1/courses/{course.id}/notebooks/{notebook_id}/document"
    monkeypatch.setattr(settings, "data_dir", tmp_path)
    monkeypatch.setattr(settings, "max_upload_bytes", 10 * 1024 * 1024)
    monkeypatch.setattr(settings, "max_pages", 100)
    monkeypatch.setattr(settings, "allowed_extensions", (".docx",))
    factory = sessionmaker(bind=db.get_bind(), autoflush=False, expire_on_commit=False)
    monkeypatch.setattr(material_worker, "SessionLocal", factory)
    monkeypatch.setattr(runner, "SessionLocal", factory)
    monkeypatch.setattr(notebook_generation_worker, "SessionLocal", factory)
    monkeypatch.setattr(runner, "claim_next_asset_job", lambda: None)
    monkeypatch.setattr(runner, "claim_next_forest_job", lambda: None)
    monkeypatch.setattr(material_worker, "extract_material_knowledge", lambda *_args: None)
    monkeypatch.setattr(material_worker, "index_material", lambda *_args: None)
    monkeypatch.setattr(material_worker, "enqueue_relations", lambda *_args, **_kwargs: None)
    generated_labels = iter(("第一资料正式内容", "第二资料正式内容", "第三资料恢复后新增内容"))
    monkeypatch.setattr(
        "app.workers.notebook_generation_worker.DeepSeekClient.complete_json",
        lambda *_args, **_kwargs: {"type": "doc", "content": [{
            "type": "paragraph", "content": [{"type": "text", "text": next(generated_labels)}],
        }]},
    )
    monkeypatch.setattr(runner.time, "monotonic", lambda: 30_000)

    def upload_parse_and_generate(title: str, body: str):
        source = BytesIO()
        document = Document()
        document.add_paragraph(body)
        document.save(source)
        response = client.post(
            f"/api/v1/notebooks/{notebook_id}/files",
            files={"file": (f"{title}.docx", source.getvalue(),
                             "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        assert response.status_code == 201, response.text
        payload = response.json()
        assert runner.claim_next_job() == payload["job"]["id"]
        material_worker.process_material(payload["job"]["id"])
        monkeypatch.setattr(runner, "_last_generation_coordination", 0)
        assert runner.run_once() is True
        db.expire_all()
        row = db.scalar(select(NotebookGenerationInput).where(
            NotebookGenerationInput.asset_id == payload["material"]["asset_id"],
        ))
        assert row.status == "completed"
        return payload

    first = upload_parse_and_generate("第一批", "第一份输入内容")
    first_revision = db.get(Notebook, notebook_id).document_revision
    second = upload_parse_and_generate("第二批", "第二份输入内容")
    assert db.get(Notebook, notebook_id).document_revision == first_revision + 1
    before_restore = client.get(document_url).json()
    assert "第一资料正式内容" in json.dumps(before_restore["content"], ensure_ascii=False)
    assert "第二资料正式内容" in json.dumps(before_restore["content"], ensure_ascii=False)

    restore = client.post(
        f"/api/v1/courses/{course.id}/notebooks/{notebook_id}/revisions/{first_revision}/restore",
        json={"expected_revision": before_restore["revision"]},
    )
    assert restore.status_code == 200, restore.text
    assert restore.json()["revision"] == before_restore["revision"] + 1
    restored_content = json.dumps(restore.json()["content"], ensure_ascii=False)
    assert "第一资料正式内容" in restored_content
    assert "第二资料正式内容" not in restored_content

    third = upload_parse_and_generate("第三批", "从未处理过的第三份输入")
    final = client.get(document_url).json()
    final_content = json.dumps(final["content"], ensure_ascii=False)
    assert final["revision"] == restore.json()["revision"] + 1
    assert "第一资料正式内容" in final_content
    assert "第二资料正式内容" not in final_content
    assert "第三资料恢复后新增内容" in final_content

    first_row = db.scalar(select(NotebookGenerationInput).where(
        NotebookGenerationInput.asset_id == first["material"]["asset_id"],
    ))
    second_row = db.scalar(select(NotebookGenerationInput).where(
        NotebookGenerationInput.asset_id == second["material"]["asset_id"],
    ))
    third_row = db.scalar(select(NotebookGenerationInput).where(
        NotebookGenerationInput.asset_id == third["material"]["asset_id"],
    ))
    assert [first_row.status, second_row.status, third_row.status] == [
        "completed", "completed", "completed",
    ]
    jobs = db.scalars(select(NotebookGenerationJob).where(
        NotebookGenerationJob.notebook_id == notebook_id,
    ).order_by(NotebookGenerationJob.created_at)).all()
    assert len(jobs) == 3 and all(job.status == "completed" for job in jobs)
    assert json.loads(jobs[-1].input_snapshot_json)["sources"][0]["asset_id"] == third["material"]["asset_id"]
    history = db.scalars(select(NotebookDocumentRevision).where(
        NotebookDocumentRevision.notebook_id == notebook_id,
    ).order_by(NotebookDocumentRevision.revision)).all()
    assert [item.origin for item in history] == ["ai_append", "ai_append", "restore", "ai_append"]
