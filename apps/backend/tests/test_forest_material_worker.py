"""Successful lecture imports automatically request forest suggestions durably."""

import json

import pytest
from docx import Document
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.db.database import Base
from app.db.models import Course, ForestJob, Material, Note, ProcessingJob
from app.workers import material_worker


@pytest.fixture
def database(monkeypatch, tmp_path):
    engine = create_engine(f"sqlite:///{(tmp_path / 'worker.sqlite3').as_posix()}")
    Base.metadata.create_all(engine)

    def factory():
        return Session(engine, autoflush=False, expire_on_commit=False)

    monkeypatch.setattr(material_worker, "SessionLocal", factory)
    monkeypatch.setattr(settings, "data_dir", tmp_path)
    for key in (
        "deepseek_api_key",
        "tavily_api_key",
        "embedding_api_key",
        "embedding_base_url",
        "embedding_model",
    ):
        monkeypatch.setattr(settings, key, None)
    settings.originals_dir.mkdir()
    with factory() as db:
        course = Course(name="线性代数")
        db.add(course)
        db.commit()
        yield db, course
    engine.dispose()


def _lecture(db, course, *, number, text):
    name = f"lecture-{number}.docx"
    document = Document()
    document.add_heading("矩阵乘法", 1)
    document.add_paragraph(text)
    document.save(settings.originals_dir / name)
    material = Material(
        course_id=course.id,
        lecture_title=f"第{number}讲",
        original_filename=name,
        stored_filename=name,
        size_bytes=100,
        chapter_order=number,
    )
    db.add(material)
    db.flush()
    job = ProcessingJob(material_id=material.id, kind="parse_material", status="processing")
    db.add(job)
    db.commit()
    return material, job


def test_first_and_subsequent_lectures_enqueue_and_retries_are_idempotent(database):
    db, course = database
    first, first_job = _lecture(db, course, number=1, text="矩阵乘法表示线性变换")
    material_worker.process_material(first_job.id)
    db.expire_all()
    assert first.status == "completed" and first_job.status == "completed"
    first_notes = list(db.scalars(select(Note).where(Note.material_id == first.id)).all())
    assert first_notes
    jobs = list(db.scalars(select(ForestJob)).all())
    assert len(jobs) == 1 and jobs[0].status == "pending" and jobs[0].course_id == course.id
    first_notes[0].content_markdown = "我的锁定笔记"
    first_notes[0].user_locked = True
    first_notes[0].revision_number += 1
    db.commit()
    _, second_job = _lecture(db, course, number=2, text="矩阵乘法还用于图形变换")
    material_worker.process_material(second_job.id)
    db.expire_all()
    assert second_job.status == "completed"
    coalesced = list(db.scalars(select(ForestJob)))
    assert len(coalesced) == 1
    expected = {note.id for note in db.scalars(select(Note).where(Note.course_id == course.id))}
    assert {row['note_id'] for row in json.loads(coalesced[0].source_snapshot_json)} == expected
    assert first_notes[0].content_markdown == "我的锁定笔记" and first_notes[0].user_locked
    repeated = ProcessingJob(
        material_id=second_job.material_id, kind="parse_material", status="processing"
    )
    db.add(repeated)
    db.commit()
    material_worker.process_material(repeated.id)
    db.expire_all()
    assert repeated.status == "completed"
    assert len(db.scalars(select(ForestJob)).all()) == 1
    assert {row['note_id'] for row in json.loads(coalesced[0].source_snapshot_json)} == expected


def test_failed_knowledge_or_index_does_not_enqueue_candidates(database, monkeypatch):
    db, course = database
    _, job = _lecture(db, course, number=1, text="矩阵乘法表示线性变换")

    def index_failed(db, material_id):
        raise RuntimeError("synthetic index failure")

    monkeypatch.setattr(material_worker, "index_material", index_failed)
    material_worker.process_material(job.id)
    db.expire_all()
    assert job.status == "failed"
    assert db.scalars(select(ForestJob)).all() == []
