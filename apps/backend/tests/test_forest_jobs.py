"""Forest imports and workers must preserve real evidence and user notes."""

import importlib
from io import BytesIO

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.core.config import settings
from app.db.database import Base, get_db
from app.db.models import Course, KnowledgeNode, Material, MaterialPage, Note, PageBlock, SourceRef


@pytest.fixture
def setup(monkeypatch, tmp_path):
    engine = create_engine(
        "sqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine)
    factory = lambda: Session(engine, autoflush=False, expire_on_commit=False)
    db = factory()
    api = FastAPI()
    try:
        routes = importlib.import_module("app.api.v1.forest_imports")
        api.include_router(routes.router, prefix="/api/v1")
        core_routes = importlib.import_module("app.api.v1.forest")
        api.include_router(core_routes.router, prefix="/api/v1")
    except ModuleNotFoundError:
        pass
    api.dependency_overrides[get_db] = lambda: db
    monkeypatch.setattr(settings, "data_dir", tmp_path)
    monkeypatch.setattr(settings, "deepseek_api_key", None)
    try:
        yield TestClient(api), db, factory
    finally:
        db.close()
        engine.dispose()


def _run(factory, monkeypatch):
    jobs = importlib.import_module("app.forest.jobs")
    monkeypatch.setattr(jobs, "SessionLocal", factory)
    claimed = jobs.claim_next_forest_job()
    assert claimed
    jobs.process_forest_job(claimed)
    return claimed


def _notes(db):
    notes = []
    for name, text in [
        ("线性代数", "矩阵乘法用于线性变换"),
        ("计算机图形学", "图形变换使用矩阵乘法"),
    ]:
        course = Course(name=name)
        db.add(course)
        db.flush()
        material = Material(
            course_id=course.id,
            lecture_title="第一讲",
            original_filename="a.pdf",
            stored_filename="a.pdf",
            size_bytes=1,
            status="completed",
        )
        node = KnowledgeNode(course_id=course.id, name="矩阵乘法")
        db.add_all([material, node])
        db.flush()
        page = MaterialPage(material_id=material.id, page_number=1, raw_text=text)
        db.add(page)
        db.flush()
        block = PageBlock(page_id=page.id, block_type="text", content=text, position=0)
        db.add(block)
        db.flush()
        ref = SourceRef(page_block_id=block.id, quote=text)
        note = Note(
            course_id=course.id,
            knowledge_node_id=node.id,
            material_id=material.id,
            title="矩阵乘法",
            content_markdown="用户保留正文",
            user_locked=True,
            revision_number=5,
        )
        db.add_all([ref, note])
        db.flush()
        note.source_refs.append(ref)
        notes.append(note)
    db.commit()
    return notes


def test_curriculum_text_is_queued_then_previewed_before_courses_exist(setup, monkeypatch):
    client, db, factory = setup
    response = client.post(
        "/api/v1/forest/curriculum-imports",
        data={"text": "第1学期 | 高等数学 | 基础\n第2学期 | 线性代数 | 基础 | 先修：高等数学"},
    )
    assert response.status_code == 201
    body = response.json()
    assert body["status"] == "pending" and body["job_id"]
    assert db.scalar(select(func.count()).select_from(Course)) == 0
    _run(factory, monkeypatch)
    preview = client.get(f"/api/v1/forest/curriculum-imports/{body['id']}").json()
    assert preview["status"] == "preview"
    assert preview["mode"] == "local_rules"
    assert preview["warnings"]
    assert [row["name"] for row in preview["rows"]] == ["高等数学", "线性代数"]
    assert preview["rows"][1]["prerequisites"] == ["高等数学"]
    assert db.scalar(select(func.count()).select_from(Material)) == 0
    rows = preview["rows"]
    rows[0]["semester"] = 3
    confirmation = client.post(
        f"/api/v1/forest/curriculum-imports/{body['id']}/confirm", json={"rows": rows}
    )
    assert confirmation.status_code == 200
    assert len(confirmation.json()["courses"]) == 2
    assert (
        client.post(
            f"/api/v1/forest/curriculum-imports/{body['id']}/confirm", json={"rows": rows}
        ).status_code
        == 409
    )


def test_confirm_rejects_forged_source_and_keeps_existing_locked_note(setup, monkeypatch):
    client, db, factory = setup
    note = _notes(db)[0]
    result = client.post(
        "/api/v1/forest/curriculum-imports", data={"text": "第1学期 | 线性代数 | 基础"}
    )
    assert result.status_code == 201
    _run(factory, monkeypatch)
    path = f"/api/v1/forest/curriculum-imports/{result.json()['id']}"
    row = client.get(path).json()["rows"][0]
    row["source_excerpt"] = "伪造的培养方案依据"
    assert client.post(path + "/confirm", json={"rows": [row]}).status_code == 422
    row = client.get(path).json()["rows"][0]
    row["existing_course_id"] = note.course_id
    assert client.post(path + "/confirm", json={"rows": [row]}).status_code == 200
    db.refresh(note)
    assert (
        note.content_markdown == "用户保留正文" and note.user_locked and note.revision_number == 5
    )
    assert db.scalar(select(func.count()).select_from(Course)) == 2


def test_docx_keeps_paragraph_table_labels_and_rejects_false_container(setup, monkeypatch):
    client, _, factory = setup
    assert (
        client.post(
            "/api/v1/forest/curriculum-imports",
            files={"file": ("fake.docx", b"fake", "application/octet-stream")},
        ).status_code
        == 415
    )
    document = Document()
    document.add_paragraph("第1学期 | 高等数学 | 基础")
    table = document.add_table(rows=1, cols=3)
    for cell, value in zip(table.rows[0].cells, ["第2学期", "计算机图形学", "专业"], strict=True):
        cell.text = value
    source = BytesIO()
    document.save(source)
    result = client.post(
        "/api/v1/forest/curriculum-imports",
        files={
            "file": (
                "plan.docx",
                source.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert result.status_code == 201
    _run(factory, monkeypatch)
    preview = client.get(f"/api/v1/forest/curriculum-imports/{result.json()['id']}").json()
    assert preview["status"] == "preview"
    assert "段落" in preview["rows"][0]["source_label"]
    assert "表格" in preview["rows"][1]["source_label"]
    assert all("页" not in row["source_label"] for row in preview["rows"])


def test_curriculum_failure_and_retry_stay_durable(setup, monkeypatch):
    client, _, factory = setup
    result = client.post("/api/v1/forest/curriculum-imports", data={"text": "培养方案"})
    assert result.status_code == 201
    path = f"/api/v1/forest/curriculum-imports/{result.json()['id']}"
    _run(factory, monkeypatch)
    assert client.get(path).json()["status"] == "failed"
    retried = client.post(path + "/retry")
    assert retried.status_code == 200
    assert retried.json()["status"] == "pending"


def test_manual_relation_requests_reuse_active_job_and_fail_without_model(setup, monkeypatch):
    client, db, factory = setup
    notes = _notes(db)
    path = f"/api/v1/forest/courses/{notes[0].course_id}/suggest-relations"
    first = client.post(path)
    assert first.status_code == 200
    assert client.post(path).json()["id"] == first.json()["id"]
    _run(factory, monkeypatch)
    status = client.get(f"/api/v1/forest/jobs/{first.json()['id']}").json()
    assert status["status"] == "failed" and "配置" in status["error_message"]
    assert client.post(f"/api/v1/forest/jobs/{first.json()['id']}/retry").status_code == 200
    notes[0].revision_number += 1
    db.commit()
    assert client.post(path).json()["id"] == first.json()["id"]


def _current_ai_relations(monkeypatch, notes):
    monkeypatch.setattr(settings, "deepseek_api_key", "test-only")

    def response(self, messages, **kwargs):
        return {
            "relations": [
                {
                    "source_note_id": notes[0].id,
                    "target_note_id": notes[1].id,
                    "source_revision": notes[0].revision_number,
                    "target_revision": notes[1].revision_number,
                    "relation_type": "application",
                    "reason": "模型推断：矩阵乘法用于图形变换",
                    "source_ref_ids": [notes[0].source_refs[0].id],
                    "target_ref_ids": [notes[1].source_refs[0].id],
                }
            ],
        }

    monkeypatch.setattr(
        importlib.import_module("app.ai.deepseek").DeepSeekClient,
        "complete_json",
        response,
    )


def test_explicit_suggestions_generate_again_when_target_notes_change(setup, monkeypatch):
    client, db, factory = setup
    notes = _notes(db)
    _current_ai_relations(monkeypatch, notes)
    path = f"/api/v1/forest/courses/{notes[0].course_id}/suggest-relations"
    first = client.post(path).json()
    _run(factory, monkeypatch)
    notes[1].revision_number += 1
    db.commit()
    second = client.post(path).json()
    assert second["id"] != first["id"] and second["status"] == "pending"
    _run(factory, monkeypatch)
    assert client.get(f"/api/v1/forest/jobs/{second['id']}").json()["status"] == "completed"
    relations = client.get("/api/v1/forest/relations?status=pending").json()
    assert len(relations) == 1 and relations[0]["target_revision"] == 6
    assert not relations[0]["stale"]


def test_explicit_suggestions_rebuild_deleted_relation_and_reuse_active_clicks(setup, monkeypatch):
    client, db, factory = setup
    notes = _notes(db)
    _current_ai_relations(monkeypatch, notes)
    path = f"/api/v1/forest/courses/{notes[0].course_id}/suggest-relations"
    first = client.post(path).json()
    _run(factory, monkeypatch)
    relation = client.get("/api/v1/forest/relations?status=pending").json()[0]
    assert client.delete(f"/api/v1/forest/relations/{relation['id']}").status_code == 200
    second = client.post(path).json()
    assert second["id"] != first["id"] and second["status"] == "pending"
    assert client.post(path).json()["id"] == second["id"]
    jobs = importlib.import_module("app.forest.jobs")
    monkeypatch.setattr(jobs, "SessionLocal", factory)
    assert jobs.claim_next_forest_job() == second["id"]
    repeated = client.post(path).json()
    assert repeated["id"] == second["id"] and repeated["status"] == "processing"
    jobs.process_forest_job(second["id"])
    restored = client.get("/api/v1/forest/relations?status=pending").json()
    assert len(restored) == 1 and restored[0]["id"] != relation["id"]
    assert restored[0]["source_note_id"] == notes[0].id
    assert restored[0]["source_ref_ids"] == [notes[0].source_refs[0].id]


def test_automatic_relation_requests_keep_completed_revision_idempotency(setup, monkeypatch):
    _, db, factory = setup
    notes = _notes(db)
    jobs = importlib.import_module("app.forest.jobs")
    automatic = jobs.enqueue_relations(db, notes[0].course_id, material_id=notes[0].material_id)
    db.commit()
    _run(factory, monkeypatch)
    same = jobs.enqueue_relations(db, notes[0].course_id, material_id=notes[0].material_id)
    assert same.id == automatic.id
    notes[0].revision_number += 1
    db.commit()
    new_revision = jobs.enqueue_relations(db, notes[0].course_id, material_id=notes[0].material_id)
    assert new_revision.id != automatic.id


def _newer_unrelated_notes(db, original_note, count=121):
    page = original_note.source_refs[0].page_block.page
    for number in range(count):
        title = f"无关知识小节 {number:03d}"
        node = KnowledgeNode(course_id=original_note.course_id, name=title)
        block = PageBlock(page_id=page.id, block_type="text", content=title, position=number + 1)
        db.add_all([node, block])
        db.flush()
        ref = SourceRef(page_block_id=block.id, quote=title)
        note = Note(
            course_id=original_note.course_id,
            material_id=original_note.material_id,
            knowledge_node_id=node.id,
            title=title,
            content_markdown=title,
        )
        db.add_all([ref, note])
        db.flush()
        note.source_refs.append(ref)
    db.commit()


@pytest.mark.parametrize("old_endpoint", ["source", "target"])
def test_old_foundation_note_survives_newer_unrelated_candidate_pool(
    setup, monkeypatch, old_endpoint
):
    client, db, factory = setup
    notes = _notes(db)
    _newer_unrelated_notes(db, notes[0])
    _current_ai_relations(monkeypatch, notes)
    course_id = notes[0].course_id if old_endpoint == "source" else notes[1].course_id
    result = client.post(f"/api/v1/forest/courses/{course_id}/suggest-relations").json()
    _run(factory, monkeypatch)
    assert client.get(f"/api/v1/forest/jobs/{result['id']}").json()["status"] == "completed"
    relations = client.get("/api/v1/forest/relations?status=pending").json()
    assert len(relations) == 1 and relations[0]["source_note_id"] == notes[0].id
    assert relations[0]["source_ref_ids"] == [notes[0].source_refs[0].id]
    assert relations[0]["target_note_id"] == notes[1].id and not relations[0]["stale"]


def test_course_latest_job_exposes_latest_automatic_failure_and_checks_course(setup, monkeypatch):
    from datetime import UTC, datetime

    client, db, factory = setup
    notes = _notes(db)
    path = f"/api/v1/forest/courses/{notes[0].course_id}/latest-job"
    assert client.get(path).status_code == 200 and client.get(path).json() is None
    assert client.get("/api/v1/forest/courses/missing/latest-job").status_code == 404
    jobs = importlib.import_module("app.forest.jobs")
    first = jobs.enqueue_relations(db, notes[0].course_id, material_id=notes[0].material_id)
    db.commit()
    _run(factory, monkeypatch)
    db.expire_all()
    failure = client.get(path).json()
    assert failure["id"] == first.id and failure["status"] == "failed"
    assert "配置" in failure["error_message"]
    models = importlib.import_module("app.db.models")
    foreign_job = models.ForestJob(
        kind="cross_course_relations",
        course_id=notes[1].course_id,
        status="failed",
        error_message="另课错误",
    )
    db.add(foreign_job)
    db.commit()
    assert client.get(path).json()["id"] == first.id
    notes[0].revision_number += 1
    db.commit()
    latest = jobs.enqueue_relations(db, notes[0].course_id, material_id=notes[0].material_id)
    db.commit()
    assert client.get(path).json()["id"] == latest.id
    notes[0].course.deleted_at = datetime.now(UTC)
    db.commit()
    assert client.get(path).status_code == 404


def test_valid_ai_candidates_pending_with_sources_and_without_note_changes(setup, monkeypatch):
    client, db, factory = setup
    notes = _notes(db)
    monkeypatch.setattr(settings, "deepseek_api_key", "test-only")
    client_type = importlib.import_module("app.ai.deepseek").DeepSeekClient
    monkeypatch.setattr(
        client_type,
        "complete_json",
        lambda self, messages, **kwargs: {
            "relations": [
                {
                    "source_note_id": notes[0].id,
                    "target_note_id": notes[1].id,
                    "source_revision": 5,
                    "target_revision": 5,
                    "relation_type": "application",
                    "reason": "模型推断：图形学将线性代数的矩阵乘法用于图形变换。",
                    "source_ref_ids": [notes[0].source_refs[0].id],
                    "target_ref_ids": [notes[1].source_refs[0].id],
                }
            ]
        },
    )
    result = client.post(f"/api/v1/forest/courses/{notes[0].course_id}/suggest-relations")
    assert result.status_code == 200
    _run(factory, monkeypatch)
    models = importlib.import_module("app.db.models")
    relation = db.scalar(select(models.CrossCourseRelation))
    assert relation is not None and relation.status == "pending" and relation.origin == "ai"
    assert relation.source_revision == 5 and relation.target_revision == 5
    for note in notes:
        db.refresh(note)
        assert (
            note.user_locked
            and note.content_markdown == "用户保留正文"
            and note.revision_number == 5
        )


@pytest.mark.parametrize("invalid", ["foreign_ref", "stale_revision", "self", "false_quote"])
def test_invalid_ai_relation_evidence_rejects_whole_batch(setup, monkeypatch, invalid):
    client, db, factory = setup
    notes = _notes(db)
    monkeypatch.setattr(settings, "deepseek_api_key", "test-only")
    relation = {
        "source_note_id": notes[0].id,
        "target_note_id": notes[1].id,
        "source_revision": 5,
        "target_revision": 5,
        "relation_type": "related",
        "reason": "模型推断：概念相关",
        "source_ref_ids": [notes[0].source_refs[0].id],
        "target_ref_ids": [notes[1].source_refs[0].id],
    }
    if invalid == "foreign_ref":
        relation["source_ref_ids"] = relation["target_ref_ids"]
    if invalid == "stale_revision":
        relation["source_revision"] = 4
    if invalid == "self":
        relation["target_note_id"] = notes[0].id
    if invalid == "false_quote":
        notes[0].source_refs[0].quote = "不在真实解析块中的引用"
        db.commit()
    monkeypatch.setattr(
        importlib.import_module("app.ai.deepseek").DeepSeekClient,
        "complete_json",
        lambda self, messages, **kwargs: {"relations": [relation]},
    )
    result = client.post(f"/api/v1/forest/courses/{notes[0].course_id}/suggest-relations")
    assert result.status_code == 200
    _run(factory, monkeypatch)
    status = client.get(f"/api/v1/forest/jobs/{result.json()['id']}").json()
    assert status["status"] == "failed"
    models = importlib.import_module("app.db.models")
    assert db.scalar(select(func.count()).select_from(models.CrossCourseRelation)) == 0


def test_claim_once_and_interrupted_forest_task_is_explicitly_retryable(setup, monkeypatch):
    client, db, factory = setup
    result = client.post(
        "/api/v1/forest/curriculum-imports", data={"text": "第1学期 | 高等数学 | 基础"}
    )
    assert result.status_code == 201
    jobs = importlib.import_module("app.forest.jobs")
    monkeypatch.setattr(jobs, "SessionLocal", factory)
    assert jobs.claim_next_forest_job() == result.json()["job_id"]
    assert jobs.claim_next_forest_job() is None
    assert jobs.mark_interrupted_forest_jobs(force=True) == 1
    models = importlib.import_module("app.db.models")
    job = db.get(models.ForestJob, result.json()["job_id"])
    db.refresh(job)
    assert job.status == "failed" and job.finished_at is not None
    assert (
        client.post(f"/api/v1/forest/curriculum-imports/{result.json()['id']}/retry").status_code
        == 200
    )


def test_relation_enqueue_is_part_of_the_callers_transaction(setup):
    _, db, _ = setup
    notes = _notes(db)
    jobs = importlib.import_module("app.forest.jobs")
    first = jobs.enqueue_relations(db, notes[0].course_id, material_id=notes[0].material_id)
    assert first is not None
    db.rollback()
    model = importlib.import_module("app.db.models").ForestJob
    assert db.scalar(select(func.count()).select_from(model)) == 0


def test_ai_curriculum_rejects_fabricated_courses_and_hides_provider_secrets(setup, monkeypatch):
    client, _, factory = setup
    monkeypatch.setattr(settings, "deepseek_api_key", "test-only")
    provider = importlib.import_module("app.ai.deepseek")
    monkeypatch.setattr(
        provider.DeepSeekClient,
        "complete_json",
        lambda self, messages, **kwargs: {
            "rows": [
                {
                    "key": "source-1",
                    "name": "不存在的课程",
                    "semester": 1,
                    "level": "foundation",
                    "prerequisites": [],
                }
            ]
        },
    )
    result = client.post(
        "/api/v1/forest/curriculum-imports", data={"text": "第1学期 | 高等数学 | 基础"}
    )
    _run(factory, monkeypatch)
    path = f"/api/v1/forest/curriculum-imports/{result.json()['id']}"
    failed = client.get(path).json()
    assert failed["status"] == "failed" and "依据" in failed["error_message"]
    assert client.post(path + "/retry").status_code == 200

    def provider_error(self, messages, **kwargs):
        raise provider.DeepSeekError("Bearer secret-must-never-be-shown")

    monkeypatch.setattr(provider.DeepSeekClient, "complete_json", provider_error)
    _run(factory, monkeypatch)
    assert "secret-must-never-be-shown" not in client.get(path).text


def test_candidate_generation_rejects_course_deleted_during_model_call(setup, monkeypatch):
    from datetime import UTC, datetime

    client, db, factory = setup
    notes = _notes(db)
    monkeypatch.setattr(settings, "deepseek_api_key", "test-only")

    def response(self, messages, **kwargs):
        notes[1].course.deleted_at = datetime.now(UTC)
        db.commit()
        return {
            "relations": [
                {
                    "source_note_id": notes[0].id,
                    "target_note_id": notes[1].id,
                    "source_revision": 5,
                    "target_revision": 5,
                    "relation_type": "application",
                    "reason": "模型推断：矩阵应用",
                    "source_ref_ids": [notes[0].source_refs[0].id],
                    "target_ref_ids": [notes[1].source_refs[0].id],
                }
            ]
        }

    monkeypatch.setattr(
        importlib.import_module("app.ai.deepseek").DeepSeekClient, "complete_json", response
    )
    result = client.post(f"/api/v1/forest/courses/{notes[0].course_id}/suggest-relations")
    _run(factory, monkeypatch)
    assert client.get(f"/api/v1/forest/jobs/{result.json()['id']}").json()["status"] == "failed"


@pytest.mark.parametrize("reviewed_status", ["confirmed", "rejected"])
def test_new_revision_refreshes_only_unreviewed_ai_candidate(setup, monkeypatch, reviewed_status):
    client, db, factory = setup
    notes = _notes(db)
    monkeypatch.setattr(settings, "deepseek_api_key", "test-only")

    def response(self, messages, **kwargs):
        return {
            "relations": [
                {
                    "source_note_id": notes[0].id,
                    "target_note_id": notes[1].id,
                    "source_revision": notes[0].revision_number,
                    "target_revision": notes[1].revision_number,
                    "relation_type": "application",
                    "reason": "模型推断：矩阵应用",
                    "source_ref_ids": [notes[0].source_refs[0].id],
                    "target_ref_ids": [notes[1].source_refs[0].id],
                }
            ]
        }

    monkeypatch.setattr(
        importlib.import_module("app.ai.deepseek").DeepSeekClient, "complete_json", response
    )
    path = f"/api/v1/forest/courses/{notes[0].course_id}/suggest-relations"
    client.post(path)
    _run(factory, monkeypatch)
    model = importlib.import_module("app.db.models").CrossCourseRelation
    relation = db.scalar(select(model))
    notes[0].revision_number = 6
    db.commit()
    client.post(path)
    _run(factory, monkeypatch)
    db.refresh(relation)
    assert relation.source_revision == 6
    relation.status = reviewed_status
    notes[0].revision_number = 7
    db.commit()
    client.post(path)
    _run(factory, monkeypatch)
    db.refresh(relation)
    assert relation.status == reviewed_status and relation.source_revision == 6


def test_curriculum_ai_calls_record_safe_telemetry(setup, monkeypatch):
    client, db, factory = setup
    monkeypatch.setattr(settings, "deepseek_api_key", "test-only")
    provider = importlib.import_module("app.ai.deepseek")
    monkeypatch.setattr(
        provider.DeepSeekClient,
        "complete_json",
        lambda self, messages, **kwargs: {
            "rows": [
                {
                    "key": "source-1",
                    "name": "线性代数",
                    "semester": 2,
                    "level": "foundation",
                    "prerequisites": [],
                }
            ]
        },
    )
    result = client.post(
        "/api/v1/forest/curriculum-imports", data={"text": "第2学期 | 线性代数 | 基础"}
    )
    _run(factory, monkeypatch)
    assert (
        client.get(f"/api/v1/forest/curriculum-imports/{result.json()['id']}").json()["mode"]
        == "ai"
    )
    audit = db.scalar(select(importlib.import_module("app.db.models").ProviderCall))
    assert audit is not None and audit.operation == "curriculum_import" and audit.duration_ms >= 0
    assert "线性代数" not in audit.metadata_json and "test-only" not in audit.metadata_json


def test_curriculum_ai_accepts_only_explicit_prerequisites(setup, monkeypatch):
    client, _, factory = setup
    monkeypatch.setattr(settings, "deepseek_api_key", "test-only")
    provider = importlib.import_module("app.ai.deepseek")
    monkeypatch.setattr(
        provider.DeepSeekClient,
        "complete_json",
        lambda self, messages, **kwargs: {
            "rows": [
                {
                    "key": "source-1",
                    "name": "线性代数",
                    "semester": 2,
                    "level": "foundation",
                    "prerequisites": ["线性代数"],
                }
            ]
        },
    )
    fabricated = client.post(
        "/api/v1/forest/curriculum-imports",
        data={"text": "第2学期 | 线性代数 | 基础 | 先修：高等数学"},
    )
    _run(factory, monkeypatch)
    assert (
        client.get(f"/api/v1/forest/curriculum-imports/{fabricated.json()['id']}").json()["status"]
        == "failed"
    )


def _curriculum_docx_table(client, headers, values):
    document = Document()
    table = document.add_table(rows=2, cols=len(headers))
    for row, content in zip(table.rows, [headers, values], strict=True):
        for cell, value in zip(row.cells, content, strict=True):
            cell.text = value
    source = BytesIO()
    document.save(source)
    result = client.post(
        "/api/v1/forest/curriculum-imports",
        files={
            "file": (
                "curriculum.docx",
                source.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert result.status_code == 201
    return f"/api/v1/forest/curriculum-imports/{result.json()['id']}"


def _ai_curriculum_row(monkeypatch, prerequisites):
    monkeypatch.setattr(settings, "deepseek_api_key", "test-only")
    monkeypatch.setattr(
        importlib.import_module("app.ai.deepseek").DeepSeekClient,
        "complete_json",
        lambda self, messages, **kwargs: {
            "rows": [
                {
                    "key": "source-2",
                    "name": "线性代数",
                    "semester": 2,
                    "level": "unspecified",
                    "prerequisites": prerequisites,
                }
            ]
        },
    )


@pytest.mark.parametrize("mode", ["local_rules", "ai"])
def test_docx_curriculum_prerequisites_are_read_from_explicit_table_header(
    setup, monkeypatch, mode
):
    client, _, factory = setup
    if mode == "ai":
        _ai_curriculum_row(monkeypatch, ["高等数学"])
    path = _curriculum_docx_table(
        client, ["课程名称", "学期", "先修课程"], ["线性代数", "2", "高等数学"]
    )
    _run(factory, monkeypatch)
    result = client.get(path).json()
    assert result["status"] == "preview" and result["mode"] == mode
    assert result["rows"][0]["semester"] == 2
    assert result["rows"][0]["prerequisites"] == ["高等数学"]
    assert result["rows"][0]["source_label"] == "表格 1 · 行 2"
    assert result["rows"][0]["source_excerpt"] == "线性代数 | 2 | 高等数学"


def test_docx_ai_cannot_use_another_column_as_prerequisite(setup, monkeypatch):
    client, _, factory = setup
    _ai_curriculum_row(monkeypatch, ["概率论"])
    path = _curriculum_docx_table(
        client, ["课程名称", "学期", "先修课程", "备注"], ["线性代数", "2", "高等数学", "概率论"]
    )
    _run(factory, monkeypatch)
    result = client.get(path).json()
    assert result["status"] == "failed" and "先修" in result["error_message"]


def test_docx_empty_prerequisite_column_does_not_shift_semester(setup, monkeypatch):
    client, _, factory = setup
    path = _curriculum_docx_table(client, ["课程名称", "先修课程", "学期"], ["线性代数", "", "2"])
    _run(factory, monkeypatch)
    result = client.get(path).json()
    assert result["status"] == "preview"
    assert result["rows"][0]["semester"] == 2 and result["rows"][0]["prerequisites"] == []
