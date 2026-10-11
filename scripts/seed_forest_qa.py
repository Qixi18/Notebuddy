"""Synthetic forest QA fixture. Never run against a user data directory."""

import json
from pathlib import Path

from docx import Document

from app.core.config import settings
from app.db.database import SessionLocal
from app.db.migrate import migrate_database
from app.db.models import (
    Course,
    ForestCourseProfile,
    KnowledgeNode,
    Material,
    MaterialPage,
    Note,
    NoteRevision,
    NoteSourceMapping,
    PageBlock,
    SourceRef,
)
from app.parsers.document import parse_document

root = Path(__file__).resolve().parents[1] / "output" / "playwright"
if settings.data_dir != root / "forest-data":
    raise SystemExit("QA requires isolated output/playwright/forest-data")
migrate_database()
with SessionLocal() as db:
    if db.query(Course).first() is not None:
        raise SystemExit("QA courses already exist; preserving them")
names = [
    "离散数学",
    "数字逻辑",
    "数据结构",
    "高等数学",
    "线性代数",
    "概率论与数理统计",
    "程序设计基础",
    "大学物理",
    "计算机组成原理",
    "操作系统",
    "计算机网络",
    "数据库系统",
    "软件工程",
    "算法设计",
    "人工智能导论",
    "机器学习",
    "编译原理",
    "数值分析",
    "信息论",
    "数字信号处理",
    "工程数学",
    "电路分析",
    "模拟电子技术",
    "嵌入式系统",
    "计算机图形学",
    "人机交互",
    "计算机视觉",
    "自然语言处理",
    "控制原理",
    "最优化方法",
    "专业英语",
    "科技写作",
    "技术伦理",
    "系统设计",
    "软件测试",
    "信息安全",
    "分布式系统",
    "云计算",
    "数据挖掘",
    "模式识别",
    "数学建模",
    "通信原理",
    "物联网",
    "机器人学",
    "工程实践",
    "应用开发",
    "项目管理",
    "需求工程",
    "形式化方法",
    "计算理论",
    "组合数学",
    "图论",
    "操作系统实践",
    "数字系统设计",
    "数据库实践",
    "智能系统实践",
    "跨学科研究",
    "课程设计",
    "综合实践",
    "毕业设计",
]
index = {"courses": [], "first_notes": []}
with SessionLocal() as db:
    for ci, name in enumerate(names):
        course = Course(name=name, description="合成浏览器验证课程")
        db.add(course)
        db.flush()
        db.add(
            ForestCourseProfile(
                course_id=course.id,
                semester=1 + ci // 10,
                level="foundation"
                if ci % 3 == 0
                else "professional"
                if ci % 3 == 1
                else "application",
                learning_state="learning" if ci < 3 else "planned",
                position=ci,
                revision=0,
                prerequisites_json="[]",
            )
        )
        index["courses"].append({"id": course.id, "name": name})
        if ci >= 3:
            continue
        for lecture in range(3):
            titles = []
            for j in range(67 if lecture < 2 else 66):
                offset = lecture * 67 + j
                titles.append(
                    ("命题逻辑" if ci == 0 else "布尔代数" if ci == 1 else "树与图")
                    if offset == 0
                    else f"{name}知识小节 {offset + 1}"
                )
            filename = f"qa-{ci}-{lecture}.docx"
            document = Document()
            for title in titles:
                document.add_heading(title, level=1)
                document.add_paragraph(
                    f"{title}的定义与运用。命题运算与布尔代数共享真值表和逻辑联结词，这是用于本地功能验证的合成材料。"
                )
            path = settings.originals_dir / filename
            document.save(path)
            parsed = parse_document(path)
            material = Material(
                course_id=course.id,
                lecture_title=f"第{lecture + 1}讲：{name}",
                original_filename=filename,
                stored_filename=filename,
                size_bytes=path.stat().st_size,
                status="completed",
                chapter_order=lecture,
                page_count=len(parsed.pages),
                parser_version=parsed.parser_version,
            )
            db.add(material)
            db.flush()
            for j, page in enumerate(parsed.pages):
                node = KnowledgeNode(course_id=course.id, name=titles[j], summary="合成资料验证")
                db.add(node)
                db.flush()
                stored_page = MaterialPage(
                    material_id=material.id,
                    page_number=page.page_number,
                    title=page.title,
                    raw_text=page.raw_text,
                    location_type=page.location_type,
                    location_label=page.location_label,
                    stable_location_key=page.stable_location_key,
                    parse_status="parsed",
                    is_active=True,
                )
                db.add(stored_page)
                db.flush()
                pblock = next(b for b in page.blocks if b.block_type == "paragraph")
                block = PageBlock(
                    page_id=stored_page.id,
                    block_type="paragraph",
                    content=pblock.content,
                    position=pblock.position,
                    object_id=pblock.object_id,
                    location_label=pblock.location_label,
                )
                db.add(block)
                db.flush()
                ref = SourceRef(
                    page_block_id=block.id,
                    quote=pblock.content,
                    status="active",
                    target_label=pblock.location_label,
                )
                db.add(ref)
                db.flush()
                body = f"# {titles[j]}\n\n{pblock.content}\n\n这是合成验证笔记，原用户资料不参与本次测试。"
                note = Note(
                    course_id=course.id,
                    knowledge_node_id=node.id,
                    material_id=material.id,
                    section_key=page.stable_location_key,
                    section_order=j + 1,
                    title=titles[j],
                    content_markdown=body,
                    content_origin="user",
                    user_locked=True,
                    revision_number=1,
                )
                note.source_refs.append(ref)
                node.source_refs.append(ref)
                db.add(note)
                db.flush()
                db.add(
                    NoteRevision(
                        note_id=note.id,
                        revision_number=1,
                        content_markdown=body,
                        content_origin="user",
                        user_locked=True,
                    )
                )
                db.add(
                    NoteSourceMapping(
                        note_id=note.id, source_ref_id=ref.id, fragment_key="whole-note"
                    )
                )
                if lecture == 0 and j == 0:
                    index["first_notes"].append(
                        {
                            "course_id": course.id,
                            "note_id": note.id,
                            "material_id": material.id,
                            "title": note.title,
                        }
                    )
    db.commit()
(root / "fixture-index.json").write_text(
    json.dumps(index, ensure_ascii=False, indent=2), encoding="utf-8"
)
print(
    json.dumps(
        {"courses": 60, "learning_courses": 3, "notes_per_learning_course": 200, "materials": 9},
        ensure_ascii=False,
    )
)
