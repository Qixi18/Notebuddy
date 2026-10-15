"""TEST ONLY: fresh, isolated notebook fixtures; never use the user's data store.

Run from apps/backend: uv run python ../../scripts/seed_knowledge_tree_qa.py
The only output is output/playwright/knowledge-tree-home/{data,fixture-index.json}.
Existing fixtures are preserved: this script refuses to reseed them.
"""

import hashlib
import json
import os
import sys
from datetime import UTC, datetime, timedelta
from pathlib import Path


def isolated_qa_dir(repo: Path) -> Path:
    expected = repo / "output/playwright/knowledge-tree-home/data"
    for parent in [expected, *expected.parents]:
        if parent.is_symlink() or (hasattr(parent, "is_junction") and parent.is_junction()):
            raise SystemExit("Test-only QA output must not redirect through a symlink or junction")
        if parent == repo:
            break
    if expected.resolve() != expected:
        raise SystemExit("Test-only QA output must stay within the fixed repository output path")
    if expected.exists() or (expected.parent / "fixture-index.json").exists():
        raise SystemExit(
            "Only a fresh fixed isolated QA directory is allowed; preserving existing files"
        )
    return expected


REPO = Path(__file__).resolve().parents[1]
# Guard BEFORE importing application settings: that import can create directories.
QA = isolated_qa_dir(REPO)
os.environ["NOTEBOOK_DATA_DIR"] = str(QA)
sys.path.insert(0, str(REPO / "apps/backend"))
from docx import Document

from app.core.config import settings
from app.db.database import SessionLocal
from app.db.migrate import migrate_database
from app.db.models import (
    Course,
    ForestCourseProfile,
    KnowledgeNode,
    Material,
    MaterialPage,
    Note,
    Notebook,
    NoteRevision,
    NoteSourceMapping,
    PageBlock,
    SourceRef,
)
from app.parsers.document import parse_document

if settings.data_dir != QA or (QA / "notebook.sqlite3").exists():
    raise SystemExit(
        "Only a fresh fixed isolated QA directory is allowed; preserving existing files"
    )
migrate_database()
now = datetime.now(UTC)
names = ["离散数学", "数字逻辑", "数学分析", "单知识点课程", "三十小节课程", "十本笔记课程"] + [
    f"培养方案课程 {i:02}" for i in range(7, 61)
]
books = [
    ["离散数学基础", "组合数学", "图论专题"],
    ["数字逻辑基础", "数字系统设计", "逻辑设计实验"],
    ["数学分析上", "数学分析下", "数学竞赛"],
]
topics = [
    ["命题逻辑", "集合与关系", "图的基本概念", "逻辑联结词", "真值表", "组合计数"],
    ["布尔代数", "逻辑门", "组合逻辑", "卡诺图", "触发器", "编码器与译码器"],
    ["函数与极限", "微分与导数", "中值定理", "Taylor展开与余项估计", "定积分", "函数列与一致收敛"],
]
index = {
    "mode": "synthetic-notes-with-real-DOCX-sources",
    "courses": [],
    "notebooks": [],
    "first_notes": [],
}
with SessionLocal() as db:
    for ci, name in enumerate(names):
        course = Course(
            name=name,
            description="隔离合成浏览器验证；正文不是模型质量证据",
            created_at=now - timedelta(days=ci),
            last_opened_at=now - timedelta(minutes=ci) if ci < 3 else None,
        )
        db.add(course)
        db.flush()
        db.add(
            ForestCourseProfile(
                course_id=course.id,
                semester=1 + ci // 10,
                level="foundation"
                if ci < 3
                else (
                    "professional"
                    if ci % 3 == 1
                    else "application"
                    if ci % 3 == 2
                    else "foundation"
                ),
                learning_state="learning" if ci < 3 else "planned",
                position=ci,
                revision=0,
                prerequisites_json="[]",
            )
        )
        index["courses"].append({"id": course.id, "name": name})
        book_names = (
            books[ci]
            if ci < 3
            else ["入门笔记"]
            if ci in (3, 4)
            else [f"专题笔记 {b + 1}" for b in range(10)]
            if ci == 5
            else []
        )
        concepts = {}
        for bi, title in enumerate(book_names):
            book = Notebook(
                course_id=course.id, title=title, tree_slot=bi, display_order=bi, is_default=bi == 0
            )
            db.add(book)
            db.flush()
            count = (
                200 if ci < 3 else 30 if ci == 4 else 1 if ci == 3 or (ci == 5 and bi == 0) else 0
            )
            index["notebooks"].append(
                {"id": book.id, "course_id": course.id, "title": title, "notes": count}
            )
            if not count:
                continue
            note_titles = [
                (topics[ci][j] if j < 6 else f"{title}知识小节 {j + 1} · 长中文与English概念")
                if ci < 3
                else f"{title}知识小节 {j + 1}"
                for j in range(count)
            ]
            filename = f"qa-{ci}-{bi}.docx"
            doc = Document()
            for heading in note_titles:
                doc.add_heading(heading, level=1)
                doc.add_paragraph(
                    f"{heading}的定义与例题。本资料用于隔离功能验证，命题运算与布尔代数共享真值表和逻辑联结词；这些文字由测试脚本编写。"
                )
            path = settings.originals_dir / filename
            doc.save(path)
            parsed = parse_document(path)
            mat = Material(
                course_id=course.id,
                notebook_id=book.id,
                lecture_title=f"第1讲：{title}",
                original_filename=filename,
                stored_filename=filename,
                size_bytes=path.stat().st_size,
                content_hash=hashlib.sha256(path.read_bytes()).hexdigest(),
                status="completed",
                chapter_order=0,
                page_count=len(parsed.pages),
                parser_version=parsed.parser_version,
            )
            db.add(mat)
            db.flush()
            for j, page in enumerate(parsed.pages):
                title = note_titles[j]
                node = concepts.get(title)
                if node is None:
                    node = KnowledgeNode(course_id=course.id, name=title, summary="合成验证概念")
                    db.add(node)
                    db.flush()
                    concepts[title] = node
                stored = MaterialPage(
                    material_id=mat.id,
                    page_number=page.page_number,
                    title=page.title,
                    raw_text=page.raw_text,
                    location_type=page.location_type,
                    location_label=page.location_label,
                    stable_location_key=page.stable_location_key,
                    parse_status="parsed",
                    is_active=True,
                )
                db.add(stored)
                db.flush()
                block = next(b for b in page.blocks if b.block_type == "paragraph")
                raw = PageBlock(
                    page_id=stored.id,
                    block_type="paragraph",
                    content=block.content,
                    position=block.position,
                    object_id=block.object_id,
                    location_label=block.location_label,
                )
                db.add(raw)
                db.flush()
                ref = SourceRef(
                    page_block_id=raw.id,
                    quote=block.content,
                    status="active",
                    target_label=block.location_label,
                )
                db.add(ref)
                db.flush()
                body = f"# {title}\n\n{block.content}\n\n这是本地合成验证笔记。\n\n" + (
                    "分段阅读与相对滚动位置验证。\n\n" * 24 if j == 0 else ""
                )
                locked = ci < 3 and bi == 0 and j < 10
                note = Note(
                    course_id=course.id,
                    notebook_id=book.id,
                    knowledge_node_id=node.id,
                    material_id=mat.id,
                    section_key=page.stable_location_key,
                    section_order=j + 1,
                    title=title,
                    content_markdown=body,
                    content_origin="user",
                    user_locked=locked,
                    revision_number=1,
                )
                note.source_refs.append(ref)
                node.source_refs.append(ref)
                db.add(note)
                db.flush()
                db.add(
                    NoteRevision(
                        note_id=note.id,
                        revision_number=1,
                        content_markdown=body,
                        content_origin="user",
                        user_locked=locked,
                    )
                )
                db.add(
                    NoteSourceMapping(
                        note_id=note.id, source_ref_id=ref.id, fragment_key="whole-note"
                    )
                )
                if j == 0:
                    index["first_notes"].append(
                        {
                            "course_id": course.id,
                            "notebook_id": book.id,
                            "note_id": note.id,
                            "material_id": mat.id,
                            "title": title,
                        }
                    )
    db.commit()
(QA.parent / "fixture-index.json").write_text(
    json.dumps(index, ensure_ascii=False, indent=2), encoding="utf-8"
)
print(
    json.dumps(
        {
            "courses": len(names),
            "notebooks": len(index["notebooks"]),
            "notes": sum(b["notes"] for b in index["notebooks"]),
            "mode": index["mode"],
        },
        ensure_ascii=False,
    )
)
